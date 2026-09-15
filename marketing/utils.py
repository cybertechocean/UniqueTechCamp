import io
import csv

def generate_excel_template():
    """
    Generate an ultra-clean, professionally styled Excel (.xlsx) template file in memory.
    Contains clear column headers, width adjustments, styling, and sample rows.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        # Graceful fallback: return a CSV template if openpyxl is not yet installed
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Full Name", "Email Address", "Subject Line", "Personalized Message"])
        writer.writerow(["Eng. David Kariuki", "david.kariuki@apexventures.co.ke", "Custom Web Architecture & AI Systems for Apex Ventures", "Hello David,\n\nWe noticed Apex Ventures is actively expanding. At UniqueTechCamp, we engineer high-performance web applications.\n\nBest regards,\nUniqueTechCamp Team"])
        return output.getvalue().encode('utf-8')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulk Email List"

    # Define Brand Styling
    header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sample_font = Font(name="Calibri", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")

    headers = ["Full Name", "Email Address", "Subject Line", "Personalized Message"]
    ws.append(headers)

    # Style Header Row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Sample Data Rows
    sample_rows = [
        [
            "Eng. David Kariuki",
            "david.kariuki@apexventures.co.ke",
            "Custom Web Architecture & AI Systems for Apex Ventures",
            "Hello David,\n\nI hope this finds you well.\n\nWe noticed Apex Ventures is actively expanding. At UniqueTechCamp, we engineer high-performance web applications integrated with 24/7 AI qualification chatbots and automated M-Pesa/Card funnels.\n\nWe would love to share a 10-minute custom architectural blueprint for your team.\n\nBest regards,\nUniqueTechCamp Solutions Team\nhttps://uniquetechcamp.org"
        ],
        [
            "Dr. Sarah Njeri",
            "dr.sarah@cityhealthgroup.com",
            "Patient Booking Automation & Portal Architecture for City Health",
            "Hi Dr. Sarah,\n\nWe recently deployed a clinical management system that reduced patient booking wait times by 40% for local healthcare providers.\n\nWe would be thrilled to show you how City Health can automate consultations and appointment reminders via WhatsApp and web.\n\nKind regards,\nUniqueTechCamp Team\n+254 715 479 955"
        ],
        [
            "Alex Munene",
            "alex@techbrandafrica.com",
            "Sub-Second High-Conversion E-Commerce Infrastructure",
            "Hi Alex,\n\nSpeed and seamless checkout are the #1 drivers of digital conversions. We specialize in sub-second load speeds and AI-driven qualification funnels.\n\nAre you available for a brief discovery chat this week?\n\nWarm regards,\nUniqueTechCamp Desk"
        ]
    ]

    for row_idx, row_data in enumerate(sample_rows, 2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 70
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = sample_font
            c.alignment = align_left
            c.border = thin_border

    # Set Column Widths
    col_widths = {
        'A': 24, # Name
        'B': 34, # Email
        'C': 40, # Subject
        'D': 65, # Message
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def parse_spreadsheet(file_obj, default_subject=""):
    """
    Parses either an Excel (.xlsx) or CSV file.
    Returns a list of dicts:
    [
      {'name': '...', 'email': '...', 'subject': '...', 'message': '...'},
      ...
    ]
    Handles header normalization (case-insensitive, strips spaces).
    """
    file_name = getattr(file_obj, 'name', '').lower()
    recipients_data = []

    if file_name.endswith('.csv'):
        # Parse CSV
        file_content = file_obj.read()
        if isinstance(file_content, bytes):
            try:
                text_content = file_content.decode('utf-8')
            except UnicodeDecodeError:
                text_content = file_content.decode('latin-1')
        else:
            text_content = file_content

        reader = csv.reader(io.StringIO(text_content))
        rows = list(reader)
        if not rows:
            return []

        # Find header indexes
        header_row = [str(h).strip().lower() for h in rows[0]]
        name_idx = _find_column_index(header_row, ['name', 'full name', 'client name', 'contact name'])
        email_idx = _find_column_index(header_row, ['email', 'email address', 'e-mail', 'mail'])
        subject_idx = _find_column_index(header_row, ['subject', 'subject line', 'title'])
        message_idx = _find_column_index(header_row, ['message', 'personalized message', 'body', 'content', 'email body'])

        if email_idx is None:
            # Fallback: check if 2nd column is email or 1st column is email
            for idx, col in enumerate(header_row):
                if '@' in col:
                    email_idx = idx
                    break
            if email_idx is None:
                email_idx = 1 if len(header_row) > 1 else 0

        for row in rows[1:]:
            if not row or not any(row):
                continue
            email = row[email_idx].strip() if email_idx < len(row) else ''
            if not email or '@' not in email:
                continue

            name = row[name_idx].strip() if name_idx is not None and name_idx < len(row) else ''
            subject = row[subject_idx].strip() if subject_idx is not None and subject_idx < len(row) else default_subject
            if not subject:
                subject = default_subject or f"Update from UniqueTechCamp for {name or 'your business'}"
            message = row[message_idx].strip() if message_idx is not None and message_idx < len(row) else ''

            recipients_data.append({
                'name': name,
                'email': email,
                'subject': subject,
                'message': message,
            })

    else:
        # Parse Excel (.xlsx) via openpyxl
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError("The 'openpyxl' Python package is required to read Excel .xlsx files. Please install openpyxl in your Python environment or upload as .csv.")
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        header_row = [str(h or '').strip().lower() for h in rows[0]]
        name_idx = _find_column_index(header_row, ['name', 'full name', 'client name', 'contact name'])
        email_idx = _find_column_index(header_row, ['email', 'email address', 'e-mail', 'mail'])
        subject_idx = _find_column_index(header_row, ['subject', 'subject line', 'title'])
        message_idx = _find_column_index(header_row, ['message', 'personalized message', 'body', 'content', 'email body'])

        if email_idx is None:
            email_idx = 1 if len(header_row) > 1 else 0

        for row in rows[1:]:
            if not row or not any(row):
                continue
            email = str(row[email_idx] or '').strip() if email_idx < len(row) else ''
            if not email or '@' not in email:
                continue

            name = str(row[name_idx] or '').strip() if name_idx is not None and name_idx < len(row) else ''
            subject = str(row[subject_idx] or '').strip() if subject_idx is not None and subject_idx < len(row) else default_subject
            if not subject:
                subject = default_subject or f"Update from UniqueTechCamp for {name or 'your business'}"
            message = str(row[message_idx] or '').strip() if message_idx is not None and message_idx < len(row) else ''

            recipients_data.append({
                'name': name,
                'email': email,
                'subject': subject,
                'message': message,
            })

    return recipients_data


def _find_column_index(header_row, possible_names):
    """Utility to locate index of a column matching any of the possible names."""
    for idx, col in enumerate(header_row):
        for name in possible_names:
            if name == col or name in col:
                return idx
    return None
